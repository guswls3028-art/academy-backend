from __future__ import annotations

import io
from dataclasses import dataclass
from types import SimpleNamespace

import pytest

from academy.application.use_cases.tools.generate_ppt import (
    _add_segmented_pdf_slides_to_composer,
    _build_pdf_question_plan,
)


@dataclass
class _Block:
    text: str
    x0: float
    y0: float
    x1: float
    y1: float


class _FakeDoc:
    def __init__(self, pages: list[list[_Block]], width: float = 600.0, height: float = 840.0):
        self._pages = pages
        self._width = width
        self._height = height

    def page_count(self) -> int:
        return len(self._pages)

    def is_blank_page(self, _page_index: int) -> bool:
        return False  # An absent text layer alone does not establish blankness.

    def extract_text_blocks(self, page_index: int) -> list[_Block]:
        return self._pages[page_index]

    def page_dimensions(self, _page_index: int) -> tuple[float, float]:
        return self._width, self._height


def _question_text(n: int) -> str:
    return (
        f"{n}. 다음 중 옳은 것은? ① 보기 하나 ② 보기 둘 ③ 보기 셋 "
        "풀이 과정과 자료를 충분히 포함한 문항 본문입니다."
    )


@pytest.mark.parametrize("middle_content", ["blank", "image", "vector", "annotation"])
def test_blank_separator_preserves_question_crops_but_visible_textless_pages_are_kept(tmp_path, middle_content):
    import fitz
    from PIL import Image
    from pptx import Presentation

    from academy.application.use_cases.tools.generate_ppt import GeneratePptFromPdfUseCase

    source = tmp_path / "questions-with-separator.pdf"
    with fitz.open() as document:
        for page_index in range(3):
            page = document.new_page(width=595, height=842)
            if page_index != 1:
                number = 1 if page_index == 0 else 2
                page.insert_text((40, 80), f"{number}. 다음 중 옳은 것은?", fontname="korea")
                page.insert_text((40, 105), "① 보기 하나 ② 보기 둘 ③ 보기 셋 ④ 보기 넷", fontname="korea")
            elif middle_content == "image":
                image = io.BytesIO()
                Image.new("RGB", (30, 30), "black").save(image, format="PNG")
                page.insert_image(fitz.Rect(40, 80, 140, 180), stream=image.getvalue())
            elif middle_content == "vector":
                page.draw_rect(fitz.Rect(40, 80, 140, 180), color=(0, 0, 0))
            elif middle_content == "annotation":
                page.add_rect_annot(fitz.Rect(40, 80, 140, 180)).update()
        document.save(source)

    result = GeneratePptFromPdfUseCase().execute(str(source))
    expected_count = 2 if middle_content == "blank" else 3
    assert result.mode == ("question" if middle_content == "blank" else "page")
    assert result.slide_count == expected_count
    presentation = Presentation(io.BytesIO(result.pptx_bytes))
    assert len(presentation.slides) == expected_count
    assert all(len(slide.shapes) == 1 for slide in presentation.slides)


def test_ppt_pdf_plan_uses_whole_page_for_scan_pdf_without_text():
    plan = _build_pdf_question_plan(_FakeDoc([[], []]))

    assert plan.use_whole_page is True
    assert plan.allow_image_segmentation is True
    assert plan.regions_per_page == [[], []]


def test_ppt_pdf_plan_preserves_image_only_pages_between_anchored_text_pages(monkeypatch):
    from academy.domain.tools import paper_type, question_splitter

    pages = [
        [_Block(f"{index}. lesson question", 40, 80, 560, 150)] if index in (1, 2, 7) else []
        for index in range(1, 8)
    ]
    monkeypatch.setattr(
        paper_type,
        "classify_paper_type",
        lambda **_kwargs: SimpleNamespace(is_non_question=False),
    )
    monkeypatch.setattr(
        question_splitter,
        "split_questions",
        lambda blocks, *_args, **_kwargs: [SimpleNamespace(number=1)] if blocks else [],
    )

    plan = _build_pdf_question_plan(_FakeDoc(pages))

    assert plan.use_whole_page is True
    assert plan.allow_image_segmentation is False
    assert plan.regions_per_page == [[] for _ in range(7)]


def test_ppt_pdf_plan_preserves_pages_when_anchors_cover_only_one_of_seven(monkeypatch):
    from academy.domain.tools import paper_type, question_splitter

    pages = [[_Block(f"page {index} lesson text", 40, 80, 560, 150)] for index in range(7)]
    monkeypatch.setattr(
        paper_type,
        "classify_paper_type",
        lambda **_kwargs: SimpleNamespace(is_non_question=False),
    )
    monkeypatch.setattr(question_splitter, "count_marginal_anchor_candidates", lambda *_args: 1)
    monkeypatch.setattr(
        question_splitter,
        "split_questions",
        lambda blocks, *_args, **_kwargs: [SimpleNamespace(number=1)]
        if "page 0" in blocks[0].text else [],
    )

    plan = _build_pdf_question_plan(_FakeDoc(pages))

    assert plan.use_whole_page is True
    assert plan.allow_image_segmentation is False
    assert plan.regions_per_page == [[] for _ in range(7)]


def test_low_coverage_pdf_result_contains_every_page_without_image_segmentation(tmp_path, monkeypatch):
    import fitz
    from pptx import Presentation

    from academy.application.use_cases.tools import generate_ppt

    source = tmp_path / "low-coverage.pdf"
    document = fitz.open()
    for number in range(3):
        page = document.new_page()
        page.insert_text((72, 72), f"Original page {number + 1}")
    document.save(source)
    document.close()

    monkeypatch.setattr(
        generate_ppt,
        "_build_pdf_question_plan",
        lambda _doc: generate_ppt._PdfQuestionPlan(
            use_whole_page=True,
            regions_per_page=[[], [], []],
            allow_image_segmentation=False,
        ),
    )
    monkeypatch.setattr(
        generate_ppt,
        "_add_segmented_pdf_slides_to_composer",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(AssertionError("segmentation must not run")),
    )

    result = generate_ppt.GeneratePptFromPdfUseCase().execute(str(source))

    assert result.mode == "page"
    assert result.slide_count == 3
    assert len(Presentation(io.BytesIO(result.pptx_bytes)).slides) == 3


def test_ppt_pdf_plan_attempts_split_for_short_text_pdf():
    pages = [[
        _Block("1. 다음 중 옳은 것은? ① ㄱ ② ㄴ", 40, 100, 560, 130),
    ]]

    plan = _build_pdf_question_plan(_FakeDoc(pages))

    assert plan.use_whole_page is False
    assert [r.number for r in plan.regions_per_page[0]] == [1]


def test_ppt_pdf_plan_applies_cross_page_anchor_validation():
    pages = [
        [
            _Block(_question_text(2), 40, 100, 560, 130),
            _Block(_question_text(3), 40, 360, 560, 390),
        ],
        [
            _Block(_question_text(3), 40, 100, 560, 130),
            _Block(_question_text(4), 40, 360, 560, 390),
        ],
    ]

    plan = _build_pdf_question_plan(_FakeDoc(pages))

    assert plan.use_whole_page is False
    assert [r.number for r in plan.regions_per_page[0]] == [2, 3]
    assert [r.number for r in plan.regions_per_page[1]] == [4]


def test_ppt_pdf_plan_prefers_marginal_anchors_for_workbook_docs():
    pages = []
    for idx in range(5):
        y = 90 + idx
        pages.append([
            _Block("1.", 24, y, 36, y + 15),
            _Block(
                "다음 자료를 읽고 물음에 답하시오. ① 보기 하나 ② 보기 둘 ③ 보기 셋",
                60,
                y,
                560,
                y + 15,
            ),
            _Block("1. 하위 항목입니다.", 90, y + 60, 560, y + 75),
            _Block("2. 하위 항목입니다.", 90, y + 85, 560, y + 100),
        ])

    plan = _build_pdf_question_plan(_FakeDoc(pages))

    assert plan.use_whole_page is False
    assert plan.workbook_doc is True
    assert [[r.number for r in page] for page in plan.regions_per_page] == [[1], [1], [1], [1], [1]]


def test_ppt_pdf_plan_preserves_short_page_restart_workbook_docs():
    pages = []
    for idx in range(3):
        y = 90 + idx
        pages.append([
            _Block("1.", 24, y, 36, y + 15),
            _Block(
                "다음 자료를 읽고 답하시오. ① 보기 하나 ② 보기 둘",
                60,
                y,
                560,
                y + 15,
            ),
            _Block("2.", 24, y + 230, 36, y + 245),
            _Block(
                "다음 설명으로 옳은 것은? ① 보기 하나 ② 보기 둘",
                60,
                y + 230,
                560,
                y + 245,
            ),
        ])

    plan = _build_pdf_question_plan(_FakeDoc(pages))

    assert plan.use_whole_page is False
    assert plan.workbook_doc is True
    assert [[r.number for r in page] for page in plan.regions_per_page] == [[1, 2], [1, 2], [1, 2]]


def test_ppt_pdf_plan_does_not_treat_short_exam_false_low_anchor_as_restart():
    pages = [
        [
            _Block(_question_text(1), 40, 90, 560, 110),
            _Block(_question_text(2), 40, 230, 560, 250),
            _Block(_question_text(3), 40, 370, 560, 390),
        ],
        [
            _Block("1. 그림 1은 보기 자료의 번호입니다.", 60, 90, 560, 110),
            _Block(_question_text(4), 40, 230, 560, 250),
            _Block(_question_text(5), 40, 370, 560, 390),
        ],
    ]

    plan = _build_pdf_question_plan(_FakeDoc(pages))

    assert plan.use_whole_page is False
    assert plan.workbook_doc is False
    assert [[r.number for r in page] for page in plan.regions_per_page] == [[1, 2, 3], [4, 5]]


def test_continuously_numbered_exam_keeps_main_questions_over_nested_steps(monkeypatch):
    from academy.domain.tools import question_splitter

    pages = []
    for page_index in range(16):
        first = page_index * 2 + 1
        page = [
            _Block(_question_text(first), 35, 44, 282, 75),
            _Block(_question_text(first + 1), 308, 44, 558, 75),
        ]
        if page_index == 3:
            page.extend([
                _Block("1. 실험 첫 번째 단계입니다.", 45, 160, 260, 177),
                _Block("2. 실험 두 번째 단계입니다.", 45, 198, 260, 215),
            ])
        pages.append(page)

    # School exam folios can resemble marginal numbers on later pages.
    monkeypatch.setattr(question_splitter, "count_marginal_anchor_candidates", lambda *_args: 1)
    plan = _build_pdf_question_plan(_FakeDoc(pages))

    assert plan.use_whole_page is False
    assert plan.workbook_doc is False
    assert [[r.number for r in page] for page in plan.regions_per_page] == [
        [index * 2 + 1, index * 2 + 2] for index in range(16)
    ]


def test_ppt_pdf_image_segmentation_fallback_adds_question_slides(tmp_path, monkeypatch):
    from PIL import Image, ImageDraw

    from academy.adapters.ai.detection import segment_dispatcher

    image_path = tmp_path / "page.png"
    image = Image.new("RGB", (320, 220), "white")
    draw = ImageDraw.Draw(image)
    draw.rectangle((20, 30, 130, 110), fill="black")
    draw.rectangle((170, 30, 290, 110), fill="black")
    image.save(image_path)

    cleanup_calls: list[list[str]] = []

    def fake_segment_questions_multipage(_pdf_path):
        return {
            "pages": [{
                "page_index": 0,
                "image_path": str(image_path),
                "boxes": [(10, 20, 140, 110), (160, 20, 140, 110)],
            }],
            "total_boxes": 2,
            "is_pdf": True,
            "tmp_dirs": ["seg-tmp"],
        }

    def fake_cleanup_pdf_seg_tmp_dirs(paths):
        cleanup_calls.append(paths)

    class _Composer:
        def __init__(self):
            self.slides: list[bytes] = []

        def add_slide(self, image_bytes: bytes):
            self.slides.append(image_bytes)

    monkeypatch.setattr(
        segment_dispatcher,
        "segment_questions_multipage",
        fake_segment_questions_multipage,
    )
    monkeypatch.setattr(
        segment_dispatcher,
        "cleanup_pdf_seg_tmp_dirs",
        fake_cleanup_pdf_seg_tmp_dirs,
    )
    composer = _Composer()

    added = _add_segmented_pdf_slides_to_composer(
        "source.pdf",
        composer=composer,
        apply_user_settings=lambda b: b,
    )

    assert added == 2
    assert len(composer.slides) == 2
    assert all(slide.startswith(b"\x89PNG") for slide in composer.slides)
    assert cleanup_calls == [["seg-tmp"]]


def test_ppt_scanned_question_shortens_only_empty_middle(tmp_path, monkeypatch):
    from io import BytesIO

    from PIL import Image, ImageDraw

    from academy.adapters.ai.detection import segment_dispatcher

    image_path = tmp_path / "scanned-question.png"
    image = Image.new("RGB", (320, 900), "white")
    draw = ImageDraw.Draw(image)
    draw.rectangle((20, 20, 120, 80), fill="black")
    draw.rectangle((20, 700, 120, 760), fill="black")
    image.save(image_path)
    monkeypatch.setattr(segment_dispatcher, "segment_questions_multipage", lambda _path: {
        "pages": [{"image_path": str(image_path), "boxes": [(0, 0, 320, 900)]}],
        "tmp_dirs": [],
    })
    monkeypatch.setattr(segment_dispatcher, "cleanup_pdf_seg_tmp_dirs", lambda _paths: None)

    class Composer:
        slides: list[bytes] = []

        def add_slide(self, image_bytes: bytes):
            self.slides.append(image_bytes)

    composer = Composer()
    assert _add_segmented_pdf_slides_to_composer(
        "source.pdf", composer=composer, apply_user_settings=lambda data: data,
    ) == 1
    with Image.open(BytesIO(composer.slides[0])) as slide:
        assert slide.width == 320
        assert slide.height < 350
        assert slide.convert("L").point(lambda pixel: 255 if pixel < 100 else 0).getbbox()


def test_ppt_pdf_shared_range_keeps_opposite_column_question_body():
    """A shared [7~8] context must not replace the actual right-column Q7."""
    from academy.domain.tools.paper_type import PaperType, PaperTypeResult
    from academy.domain.tools.question_splitter import TextBlock, split_questions

    page_width, page_height = 595.0, 842.0
    paper_type = PaperTypeResult(
        paper_type=PaperType.CLEAN_PDF_DUAL,
        confidence=0.99,
        is_dual_column=True,
        is_quadrant=False,
        is_handwriting_present=False,
        has_embedded_text=True,
        debug={
            "has_embedded_text": True,
            "is_dual_text": True,
            "is_dual_pixel": False,
        },
    )
    blocks = [
        TextBlock(
            text="[7~8] 그림은 수소 원자의 중심 입자 A와 주위를 운동하는 입자 B이다.",
            x0=35.0,
            y0=44.0,
            x1=275.0,
            y1=79.0,
        ),
        TextBlock(
            text="문항 7~8 원자 모형 그림 삽입 위치",
            x0=37.0,
            y0=340.0,
            x1=55.0,
            y1=567.0,
        ),
        TextBlock(
            text="7. A와 B를 옳게 짝 지은 것은?",
            x0=308.0,
            y0=44.0,
            x1=494.0,
            y1=61.0,
        ),
        TextBlock(
            text="A B ① 양성자 중성자 ② 중성자 전자 ③ 양성자 전자",
            x0=338.0,
            y0=104.0,
            x1=519.0,
            y1=247.0,
        ),
    ]

    regions = split_questions(
        blocks,
        page_width,
        page_height,
        page_index=4,
        paper_type=paper_type,
        prefer_marginal=True,
    )

    assert [region.number for region in regions] == [7]
    region = regions[0]
    assert region.bbox[0] < page_width * 0.10
    assert region.bbox[2] > page_width * 0.85
    assert region.bbox[3] > page_height * 0.65
    assert "shared_context_cross_column" in region.semantic_flags
